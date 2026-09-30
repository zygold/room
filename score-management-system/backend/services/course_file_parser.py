"""课程设置文件解析 — 支持 Excel(.xlsx/.xlsm/.xls) / Word(.docx) / PDF.

统一输出: [{class_name, course_name, teacher_name, periods, sheet}]
文件表格需包含「班级」与「课程名称」两列（表头支持常见别名）。
"""
import io
import os
import re

_HEADER_ALIASES = {
    "class_name": ("班级", "班级名称", "行政班", "教学班", "班别"),
    "course_name": ("课程名称", "课程名", "课程", "科目名称", "科目"),
    "teacher_name": ("任课教师", "授课教师", "上课教师", "教师姓名", "教师"),
    "periods": ("节数", "课时数", "周课时", "课时", "总课时"),
}

MAX_HEADER_SCAN = 10


def _norm(value):
    if value is None:
        return ""
    return str(value).strip()


def _match_header(text):
    t = _norm(text)
    if not t or len(t) > 8:
        return None
    for key, names in _HEADER_ALIASES.items():
        if t in names:
            return key
    for key, names in _HEADER_ALIASES.items():
        for name in names:
            if name in t:
                return key
    return None


def _locate_header(grid):
    for i in range(min(len(grid), MAX_HEADER_SCAN)):
        mapping = {}
        for j, cell in enumerate(grid[i]):
            key = _match_header(cell)
            if key and key not in mapping:
                mapping[key] = j
        if "class_name" in mapping and "course_name" in mapping:
            return i, mapping
    return -1, {}


def _to_int(value):
    if value in (None, ""):
        return None
    m = re.search(r"\d+", str(value))
    return int(m.group()) if m else None


def _rows_from_grid(grid, sheet):
    if not grid:
        return []
    header_row, cols = _locate_header(grid)
    if header_row < 0:
        return []
    out, last_class = [], ""
    for row in grid[header_row + 1:]:

        def cell(key):
            j = cols.get(key)
            if j is None or j >= len(row):
                return ""
            return _norm(row[j])

        class_name = cell("class_name") or last_class
        course_name = cell("course_name")
        if not class_name or not course_name:
            continue
        last_class = class_name
        out.append({
            "class_name": class_name,
            "course_name": course_name,
            "teacher_name": cell("teacher_name"),
            "periods": _to_int(cell("periods")),
            "sheet": sheet,
        })
    return out


def _grids_xlsx(data):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True)
    grids = []
    for ws in wb.worksheets:
        rows = [[c.value for c in row] for row in ws.iter_rows()]
        # 合并单元格: 用左上角的值回填整个区域
        for rng in ws.merged_cells.ranges:
            value = ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    if 1 <= r <= len(rows) and 1 <= c <= len(rows[r - 1]):
                        rows[r - 1][c - 1] = value
        grids.append((ws.title, rows))
    return grids


def _grids_xls(data):
    from python_calamine import CalamineWorkbook
    wb = CalamineWorkbook.from_filelike(io.BytesIO(data))
    grids = []
    for name in wb.sheet_names:
        grids.append((name, wb.get_sheet_by_name(name).to_python()))
    return grids


def _grids_docx(data):
    from docx import Document
    doc = Document(io.BytesIO(data))
    grids = []
    for i, table in enumerate(doc.tables):
        grid = [[cell.text for cell in row.cells] for row in table.rows]
        grids.append(("表格%d" % (i + 1), grid))
    return grids


def _grids_pdf(data):
    import pdfplumber
    grids = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for i, page in enumerate(pdf.pages):
            for j, table in enumerate(page.extract_tables() or []):
                grid = [["" if c is None else c for c in row] for row in table]
                grids.append(("第%d页-表%d" % (i + 1, j + 1), grid))
    return grids


def _dedupe(rows):
    """同一 (班级, 课程名称) 只保留一条，教师/节数取有效值。"""
    merged = {}
    for r in rows:
        key = (r["class_name"], r["course_name"])
        cur = merged.get(key)
        if cur is None:
            merged[key] = dict(r)
            continue
        if not cur.get("teacher_name") and r.get("teacher_name"):
            cur["teacher_name"] = r["teacher_name"]
        if r.get("periods") and (not cur.get("periods") or r["periods"] > cur["periods"]):
            cur["periods"] = r["periods"]
    return list(merged.values())


def parse_course_file(filename, data):
    """解析课程设置文件，返回去重后的行列表。"""
    ext = os.path.splitext(filename or "")[1].lower()
    if ext in (".xlsx", ".xlsm"):
        grids = _grids_xlsx(data)
    elif ext == ".xls":
        grids = _grids_xls(data)
    elif ext == ".docx":
        grids = _grids_docx(data)
    elif ext == ".pdf":
        grids = _grids_pdf(data)
    else:
        raise ValueError("不支持的文件格式 %s（仅支持 Excel / Word / PDF）" % (ext or "未知"))

    rows = []
    for sheet, grid in grids:
        rows.extend(_rows_from_grid(grid, sheet))
    return _dedupe(rows)
