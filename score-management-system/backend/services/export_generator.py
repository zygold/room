"""Export generator for scores and scholarship data."""
import io
import os
import zipfile
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet

from config import EXPORT_DIR
from database import get_db
from repositories.base import BaseRepository

_base_repo = BaseRepository()

EXPORT_DIR.mkdir(parents=True, exist_ok=True)


def build_query(payload: dict):
    """Build SQL and params for score export."""
    sql = """
    SELECT s.*, st.name as student_name, st.homeroom_teacher,
           c.name as class_name, g.name as grade_name, m.name as major_name, ct.name as class_type_name,
           e.name as exam_name, e.exam_type, e.school_year, e.semester, e.month as exam_month
    FROM scores s
    JOIN students st ON s.student_id=st.id
    JOIN classes c ON st.class_id=c.id
    JOIN grades g ON st.grade_id=g.id
    JOIN majors m ON st.major_id=m.id
    JOIN class_types ct ON st.class_type_id=ct.id
    JOIN exams e ON s.exam_id=e.id
    WHERE 1=1
    """
    params = []
    if payload.get("grade_id"):
        sql += " AND st.grade_id=?"
        params.append(payload["grade_id"])
    if payload.get("major_id"):
        sql += " AND st.major_id=?"
        params.append(payload["major_id"])
    if payload.get("class_type_id"):
        sql += " AND st.class_type_id=?"
        params.append(payload["class_type_id"])
    if payload.get("class_id"):
        sql += " AND st.class_id=?"
        params.append(payload["class_id"])
    if payload.get("exam_id"):
        sql += " AND s.exam_id=?"
        params.append(payload["exam_id"])
    exam_ids = payload.get("exam_ids")
    if exam_ids:
        placeholders = ",".join(["?"] * len(exam_ids))
        sql += f" AND s.exam_id IN ({placeholders})"
        params.extend(exam_ids)
    if payload.get("exam_type"):
        sql += " AND e.exam_type=?"
        params.append(payload["exam_type"])
    if payload.get("school_year"):
        sql += " AND e.school_year=?"
        params.append(payload["school_year"])
    if payload.get("semester"):
        sql += " AND e.semester=?"
        params.append(payload["semester"])
    if payload.get("month"):
        sql += " AND e.month=?"
        params.append(payload["month"])
    if payload.get("keyword"):
        sql += " AND st.name LIKE ?"
        params.append(f"%{payload['keyword']}%")
    if payload.get("selected_ids"):
        placeholders = ",".join(["?"] * len(payload["selected_ids"]))
        sql += f" AND s.id IN ({placeholders})"
        params.extend(payload["selected_ids"])
    sql += " ORDER BY st.grade_id, st.class_id, s.total_score DESC"
    return sql, params


def fetch_data(payload: dict):
    with get_db() as conn:
        sql, params = build_query(payload)
        return _base_repo.query(sql, params, conn=conn)


def to_xlsx(rows: list, desensitize: bool = False, include_converted: bool = False):
    wb = Workbook()
    ws = wb.active
    ws.title = "成绩明细"

    if desensitize:
        headers = ["姓名", "年级", "专业", "班级", "班级类别", "考试", "考试类型",
                   "语文", "数学", "英语", "专业课", "总分"]
        ws.append(headers)
        for r in rows:
            ws.append([
                r["student_name"], r["grade_name"], r["major_name"], r["class_name"], r["class_type_name"],
                r["exam_name"], r["exam_type"],
                r["chinese_score"], r["math_score"], r["english_score"],
                r["professional_score"], r["total_score"]
            ])
        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)
        return bio.getvalue()

    # Group by exam and student: one row per student, columns per exam.
    exams = []
    seen_exam = set()
    for r in rows:
        key = (r["exam_id"], r["exam_name"])
        if key not in seen_exam:
            seen_exam.add(key)
            exams.append({"id": r["exam_id"], "name": r["exam_name"], "type": r["exam_type"]})

    students = {}
    for r in rows:
        sid = r["student_id"]
        if sid not in students:
            students[sid] = {
                "student_name": r["student_name"],
                "grade_name": r["grade_name"],
                "major_name": r["major_name"],
                "class_name": r["class_name"],
                "class_type_name": r["class_type_name"],
                "scores": {},
            }
        students[sid]["scores"][r["exam_id"]] = r

    base_headers = ["姓名", "年级", "专业", "班级类别", "班级"]
    exam_headers = []
    for exam in exams:
        prefix = exam["name"]
        if include_converted:
            exam_headers.extend([
                f"{prefix}_语文", f"{prefix}_语文(换算)",
                f"{prefix}_数学", f"{prefix}_数学(换算)",
                f"{prefix}_英语", f"{prefix}_英语(换算)",
                f"{prefix}_专业课", f"{prefix}_专业课(换算)",
                f"{prefix}_总分",
            ])
        else:
            exam_headers.extend([
                f"{prefix}_语文",
                f"{prefix}_数学",
                f"{prefix}_英语",
                f"{prefix}_专业课",
                f"{prefix}_总分",
            ])
    ws.append(base_headers + exam_headers)

    for sid in sorted(students.keys(), key=lambda x: (students[x]["grade_name"] or "", students[x]["class_name"] or "", students[x]["student_name"] or "")):
        stu = students[sid]
        row = [
            stu["student_name"], stu["grade_name"], stu["major_name"],
            stu["class_type_name"], stu["class_name"]
        ]
        for exam in exams:
            r = stu["scores"].get(exam["id"])
            if r:
                if include_converted:
                    row.extend([
                        r["chinese_score"], r["chinese_converted"],
                        r["math_score"], r["math_converted"],
                        r["english_score"], r["english_converted"],
                        r["professional_score"], r["professional_converted"],
                        r["total_score"],
                    ])
                else:
                    row.extend([
                        r["chinese_score"],
                        r["math_score"],
                        r["english_score"],
                        r["professional_score"],
                        r["total_score"],
                    ])
            else:
                row.extend([""] * (9 if include_converted else 5))
        ws.append(row)

    # Add a summary sheet listing exams.
    if exams:
        summary = wb.create_sheet(title="考试列表")
        summary.append(["考试名称", "考试类型"])
        for exam in exams:
            summary.append([exam["name"], exam["type"]])

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio.getvalue()


def to_docx(rows: list, desensitize: bool = False):
    doc = Document()
    doc.add_heading("成绩明细", level=1)
    table = doc.add_table(rows=1, cols=10 if not desensitize else 6)
    hdr_cells = table.rows[0].cells
    headers = ["姓名", "年级", "专业", "班级", "班级类别", "考试", "考试类型", "语文", "数学", "总分"]
    if desensitize:
        headers = headers[:6]
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
    for r in rows[:1000]:
        row_cells = table.add_row().cells
        vals = [r["student_name"], r["grade_name"], r["major_name"], r["class_name"], r["class_type_name"],
                r["exam_name"], r["exam_type"], str(r["chinese_score"] or ""), str(r["math_score"] or ""), str(r["total_score"] or "")]
        if desensitize:
            vals = vals[:6]
        for i, v in enumerate(vals):
            row_cells[i].text = str(v)
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio.getvalue()


def to_pdf(rows: list, desensitize: bool = False):
    bio = io.BytesIO()
    doc = SimpleDocTemplate(bio, pagesize=A4)
    elements = []
    styles = getSampleStyleSheet()
    elements.append(Paragraph("成绩明细", styles["Title"]))
    data = [["姓名", "年级", "专业", "班级", "班级类别", "考试", "总分"]]
    for r in rows[:500]:
        data.append([
            r["student_name"], r["grade_name"], r["major_name"], r["class_name"],
            r["class_type_name"], r["exam_name"], str(r["total_score"] or "")
        ])
    t = Table(data)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4f46e5')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 1, colors.grey),
    ]))
    elements.append(t)
    doc.build(elements)
    bio.seek(0)
    return bio.getvalue()


def to_md(rows: list):
    lines = ["# 成绩明细\n"]
    lines.append("| 姓名 | 年级 | 专业 | 班级 | 考试 | 总分 |")
    lines.append("|------|------|------|------|------|------|")
    for r in rows:
        lines.append(f"| {r['student_name']} | {r['grade_name']} | {r['major_name']} | {r['class_name']} | {r['exam_name']} | {r['total_score']} |")
    return "\n".join(lines).encode("utf-8")


def generate(payload: dict, fmt: str, desensitize: bool = False, options: dict = None):
    rows = fetch_data(payload)
    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    basename = f"export_{timestamp}"
    options = options or {}
    include_converted = options.get("include_converted", False)
    split_by_class = options.get("split_by_class", False)

    if split_by_class and fmt in ("xlsx", "zip"):
        data = to_xlsx_split_by_class(rows, desensitize, include_converted, basename)
        filename = f"{basename}_按班级.zip"
    elif fmt == "xlsx":
        data = to_xlsx(rows, desensitize, include_converted)
        filename = f"{basename}.xlsx"
    elif fmt == "docx":
        data = to_docx(rows, desensitize)
        filename = f"{basename}.docx"
    elif fmt == "pdf":
        data = to_pdf(rows, desensitize)
        filename = f"{basename}.pdf"
    elif fmt == "md":
        data = to_md(rows)
        filename = f"{basename}.md"
    elif fmt == "zip":
        bio = io.BytesIO()
        with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr(f"{basename}.xlsx", to_xlsx(rows, desensitize, include_converted))
            zf.writestr(f"{basename}.md", to_md(rows))
        bio.seek(0)
        data = bio.getvalue()
        filename = f"{basename}.zip"
    else:
        raise ValueError("不支持的格式")

    path = EXPORT_DIR / filename
    path.write_bytes(data)
    return path, len(data)


def to_xlsx_split_by_class(rows: list, desensitize: bool = False, include_converted: bool = False, basename: str = "export"):
    """Generate a zip archive with one xlsx file per class."""
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as zf:
        by_class = {}
        for r in rows:
            by_class.setdefault(r["class_name"], []).append(r)
        for class_name, class_rows in sorted(by_class.items()):
            safe_name = "".join(c for c in class_name if c.isalnum() or c in "_-")
            xlsx_bytes = to_xlsx(class_rows, desensitize, include_converted)
            zf.writestr(f"{basename}_{safe_name}.xlsx", xlsx_bytes)
    bio.seek(0)
    return bio.getvalue()
