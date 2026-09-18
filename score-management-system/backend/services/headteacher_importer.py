"""班主任 Excel 批量导入：兼容多种 Excel 排版.

移植自 JX 桌面版 business/headteacher_importer.py，Web 版最小改动.
"""
import re
import tempfile
import os
from openpyxl import load_workbook

_HEAD_RE = re.compile(r"班主任[：:]?\s*(.+)")


def _looks_like_class(text):
    t = str(text or "").strip()
    return "班" in t and len(t) <= 30


def _looks_like_header(row):
    texts = [str(c or "").strip() for c in row]
    has_class = any("班级" in t or t == "班" for t in texts)
    has_head = any("班主任" in t or "教师" in t or t == "姓名" for t in texts)
    return has_class and has_head


def _find_header_row(rows):
    for i, r in enumerate(rows[:20]):
        if _looks_like_header(r):
            return i
    return None


def _cell_text(cell):
    if cell is None:
        return ""
    t = str(cell).strip()
    return t if t != "None" else ""


def _resolve_columns(header):
    class_idx = head_idx = None
    for i, c in enumerate(header):
        text = _cell_text(c)
        if class_idx is None and ("班级" in text or text == "班"):
            class_idx = i
        if head_idx is None and ("班主任" in text or "教师" in text or text == "姓名"):
            head_idx = i
    if class_idx is None and len(header) >= 2:
        class_idx = 0
    if head_idx is None and len(header) >= 2:
        head_idx = 1 if class_idx == 0 else 0
    return class_idx, head_idx


def _extract_from_cell(text):
    t = str(text or "").strip()
    m = _HEAD_RE.search(t)
    if not m:
        return None
    head = m.group(1).strip()
    cls = t[:m.start()].strip()
    if not cls:
        return None
    return cls, head


def _sniff_two_column_rows(rows):
    if not rows:
        return []
    max_cols = max(len(r) for r in rows)
    if max_cols < 2:
        return []
    best = []
    for ci in range(max_cols):
        for hi in range(max_cols):
            if ci == hi:
                continue
            class_like = sum(1 for r in rows if len(r) > ci and _looks_like_class(r[ci]))
            head_like = sum(1 for r in rows if len(r) > hi
                            and not _looks_like_class(r[hi]) and _cell_text(r[hi]))
            best.append((class_like + head_like, ci, hi))
    if not best:
        return []
    best.sort(reverse=True)
    _, ci, hi = best[0]
    return [(r[ci], r[hi]) for r in rows if len(r) > max(ci, hi)
            and _cell_text(r[ci]) and _cell_text(r[hi])]


def _parse_sheet(ws):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if not rows:
        return []
    header_idx = _find_header_row(rows)
    if header_idx is not None:
        header = rows[header_idx]
        class_idx, head_idx = _resolve_columns(header)
        if class_idx is not None and head_idx is not None:
            result = []
            for r in rows[header_idx + 1:]:
                if len(r) <= max(class_idx, head_idx):
                    continue
                cls = _cell_text(r[class_idx])
                head = _cell_text(r[head_idx])
                if cls and head:
                    result.append((cls, head))
            return result
    inferred = _sniff_two_column_rows(rows)
    if inferred:
        return inferred
    result = []
    for r in rows:
        for cell in r:
            hit = _extract_from_cell(cell)
            if hit:
                result.append(hit)
    return result


def parse_headteacher_file(file_bytes):
    """Web 版入口: 直接吃 bytes, 返回 [(class_raw, head_raw), ...]."""
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp.write(file_bytes)
        tmp_path = tmp.name
    try:
        wb = load_workbook(tmp_path, read_only=True, data_only=True)
        result = []
        for ws in wb.worksheets:
            result.extend(_parse_sheet(ws))
        wb.close()
        return result
    finally:
        os.unlink(tmp_path)


def resolve_and_prepare(conn, pairs):
    """把 (class_raw, head_raw) 解析成可执行的 UPDATE 批次.

    返回 {"matched": [...], "unmatched_class": [...]}
    matched 每项: (class_id, class_name, head_name, class_raw, head_raw)
    """
    class_basic = _classes_repo.list_basic(conn=conn)
    class_registry = {name: cid for cid, name in class_basic}
    class_list = [{"id": cid, "name": name} for cid, name in class_basic]

    matched = []
    unmatched_class = []

    for cls_raw, head_raw in pairs:
        cls_raw = str(cls_raw or "").strip()
        head_raw = str(head_raw or "").strip()
        if not cls_raw or not head_raw:
            continue
        cid = class_registry.get(cls_raw)
        cname = cls_raw
        if cid is None:
            clean = re.sub(r"[（(].*?[)）]", "", cls_raw).strip()
            cid = class_registry.get(clean)
            cname = clean
        if cid is None:
            for r in class_list:
                if cls_raw in r["name"] or r["name"] in cls_raw:
                    cid = r["id"]
                    cname = r["name"]
                    break
        if cid is None:
            unmatched_class.append((cls_raw, head_raw))
        else:
            matched.append((cid, cname, head_raw, cls_raw, head_raw))

    return {"matched": matched, "unmatched_class": unmatched_class}


def apply_batch(conn, matched):
    """批量 UPDATE classes.head_teacher + students.homeroom_teacher. 返回应用数量."""
    applied = 0
    for cid, _cname, head, _cr, _hr in matched:
        _classes_repo.update_head_teacher_batch(cid, head, conn=conn)
        _students_repo.update_homeroom_by_class(cid, head, conn=conn)
        applied += 1
    return applied
